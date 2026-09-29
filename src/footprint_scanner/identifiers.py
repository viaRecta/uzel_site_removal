"""Client identifiers, read from the tracker's Client Profile sheet or a client.yaml file."""

from __future__ import annotations

import datetime as dt
import logging
import re
import warnings
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
import yaml

from .textnorm import fold

log = logging.getLogger(__name__)

PROFILE_SHEET = "Client Profile"

# Field name -> folded label prefix on the Client Profile sheet (column A).
LABELS: dict[str, str] = {
    "case_id": "client id",
    "full_name": "full legal name",
    "name_variants": "name variants",
    "dob": "date of birth",
    "cities": "current & past cities",
    "addresses": "current & past addresses",
    "phones": "phone numbers",
    "emails": "email addresses",
    "usernames": "usernames",
    "employers": "employers",
    "schools": "schools",
    "relatives": "known relatives",
    "authorized": "signed authorization received",
    "keep_accounts": "accounts client wants to keep",
}
LIST_FIELDS = (
    "name_variants", "cities", "addresses", "phones", "emails", "usernames",
    "employers", "schools", "relatives", "keep_accounts",
)

# Template placeholders such as '+90 5xx xxx xx xx', 'xxx@...', '<name>', 'TBD'.
_PLACEHOLDER = re.compile(r"\dx|x\d|\bx{2,}\b|\.\.\.|<[^>]+>|^tbd$|^n/?a$|^-$", re.IGNORECASE)


@dataclass
class ClientIdentifiers:
    case_id: str
    full_name: str
    name_variants: list[str] = field(default_factory=list)
    dob: dt.date | None = None
    cities: list[str] = field(default_factory=list)
    addresses: list[str] = field(default_factory=list)
    phones: list[str] = field(default_factory=list)
    emails: list[str] = field(default_factory=list)
    usernames: list[str] = field(default_factory=list)
    employers: list[str] = field(default_factory=list)
    schools: list[str] = field(default_factory=list)
    relatives: list[str] = field(default_factory=list)
    keep_accounts: list[str] = field(default_factory=list)
    authorized: bool = False
    source: str = ""

    @property
    def names(self) -> list[str]:
        """Full name first, then variants, deduplicated after folding."""
        seen: set[str] = set()
        out: list[str] = []
        for n in [self.full_name, *self.name_variants]:
            key = fold(n).strip()
            if key and key not in seen:
                seen.add(key)
                out.append(n.strip())
        return out

    @property
    def usernames_clean(self) -> list[str]:
        return [u.strip().lstrip("@") for u in self.usernames if u.strip().lstrip("@")]


def split_values(raw: Any) -> list[str]:
    if raw is None:
        return []
    if isinstance(raw, (list, tuple)):
        items = [str(x) for x in raw]
    else:
        items = str(raw).split(";")
    return [i.strip() for i in items if i and i.strip()]


def _is_placeholder(value: str) -> bool:
    return bool(_PLACEHOLDER.search(value.strip()))


def _valid(field_name: str, value: str) -> bool:
    if _is_placeholder(value):
        return False
    if field_name == "phones":
        return len(re.sub(r"\D", "", value)) >= 7
    if field_name == "emails":
        return "@" in value and "." in value.split("@")[-1]
    return True


def _parse_bool(raw: Any) -> bool:
    return fold(str(raw or "")).strip() in {"yes", "y", "true", "1", "evet"}


def _parse_date(raw: Any) -> dt.date | None:
    if raw in (None, ""):
        return None
    if isinstance(raw, dt.datetime):
        return raw.date()
    if isinstance(raw, dt.date):
        return raw
    s = str(raw).strip()
    for f in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            continue
    log.warning("Could not parse date of birth; ignoring it")
    return None


def _build(values: dict[str, Any], source: str) -> ClientIdentifiers:
    case_id = str(values.get("case_id") or "").strip()
    full_name = str(values.get("full_name") or "").strip()
    if not case_id:
        raise ValueError("Client ID / Case # is missing.")
    if not full_name or _is_placeholder(full_name):
        raise ValueError("Full legal name is missing.")
    kwargs: dict[str, Any] = {}
    for f in LIST_FIELDS:
        items = split_values(values.get(f))
        kept = [v for v in items if _valid(f, v)]
        if len(kept) != len(items):
            log.warning("Skipped %d placeholder/invalid value(s) in %s", len(items) - len(kept), f)
        kwargs[f] = kept
    return ClientIdentifiers(
        case_id=case_id,
        full_name=full_name,
        dob=_parse_date(values.get("dob")),
        authorized=_parse_bool(values.get("authorized")),
        source=source,
        **kwargs,
    )


def read_profile_sheet(ws) -> dict[str, Any]:
    """Map label text in column A to the value in column B."""
    values: dict[str, Any] = {}
    for row in ws.iter_rows(min_col=1, max_col=2):
        label, value = row[0].value, row[1].value
        if not isinstance(label, str):
            continue
        flabel = fold(label).strip()
        for fname, prefix in LABELS.items():
            if fname not in values and flabel.startswith(prefix):
                values[fname] = value
                break
    return values


def load_from_tracker(path: str | Path) -> ClientIdentifiers:
    with warnings.catch_warnings():
        # Excel's extended dropdown format isn't needed to read identifiers.
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if PROFILE_SHEET not in wb.sheetnames:
            raise ValueError(f"Sheet '{PROFILE_SHEET}' not found in {path}")
        values = read_profile_sheet(wb[PROFILE_SHEET])
    finally:
        wb.close()
    return _build(values, source=f"tracker:{Path(path).name}")


def load_from_yaml(path: str | Path) -> ClientIdentifiers:
    with Path(path).open(encoding="utf-8") as fh:
        data = yaml.safe_load(fh) or {}
    unknown = set(data) - set(LABELS)
    if unknown:
        log.warning("Unknown client.yaml fields ignored: %s", ", ".join(sorted(unknown)))
    return _build(data, source=f"yaml:{Path(path).name}")
