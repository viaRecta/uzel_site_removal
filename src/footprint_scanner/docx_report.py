"""Customer report as a Word document (.docx), meant to be reviewed/edited in Word and printed
or saved as PDF from there. Same content as the PDF report: findings with screenshots, a
Remove / Keep tick box per item, and a signed authorization page. Sends nothing anywhere.
"""

from __future__ import annotations

import datetime as dt
import logging
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .excel_io import (
    H_CATEGORY,
    H_CONFIDENCE,
    H_CONTACT,
    H_CONTENT,
    H_DATA_TYPES,
    H_DATE_FOUND,
    H_DOMAIN,
    H_FALLBACK,
    H_JURISDICTION,
    H_MENTIONS,
    H_METHOD,
    H_OWN,
    H_OWNER,
    H_PLACES,
    H_SEVERITY,
    H_STATUS,
    H_URL,
    Tracker,
)
from .identifiers import ClientIdentifiers
from .pdf_report import (
    SEVERITY_ORDER,
    TEXT,
    ReportItem,
    ReportMeta,
    _fmt_date,
    screenshot_jpeg,
    select_items,
    tr,
    tr_list,
)

log = logging.getLogger(__name__)

FONT = "Arial"
SYMBOL_FONT = "Segoe UI Symbol"  # has the ☐ glyph; Word substitutes a similar font elsewhere
BOX = "☐"

INK = RGBColor(0x1D, 0x29, 0x39)
MUTED = RGBColor(0x47, 0x54, 0x67)
LINK = RGBColor(0x18, 0x49, 0xA9)
SEVERITY_HEX = {"Critical": "B42318", "High": "C4320A", "Medium": "B54708", "Low": "475467"}

PAGE_W, MARGIN = 210, 18  # mm, A4
CONTENT_W = PAGE_W - 2 * MARGIN

# Print-specific wording (the PDF version talks about on-screen form fields instead).
PRINT_TEXT = {
    "en": {
        "intro3": "Please tick <b>Remove</b> or <b>Keep</b> on each item page, sign the last page and return "
                  "all pages to us. Screenshots show each page as it looked when we checked it.",
    },
    "tr": {
        "intro3": "Lütfen her madde sayfasında <b>Kaldır</b> veya <b>Kalsın</b> kutusunu işaretleyin, son sayfayı "
                  "imzalayın ve tüm sayfaları bize geri iletin. Ekran görüntüleri, sayfaların kontrol ettiğimiz "
                  "andaki halini gösterir.",
    },
}


# ---------------------------------------------------------------- low-level helpers


def _set_run_font(run, name: str = FONT, size: float | None = None, bold: bool | None = None,
                  color: RGBColor | None = None) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _add_rich(paragraph, text: str, size: float = 10, color: RGBColor = INK, bold: bool = False) -> None:
    """Add text that may contain <b>…</b> markup (shared with the PDF strings)."""
    parts = re.split(r"(<b>|</b>)", text)
    is_bold = bold
    for part in parts:
        if part == "<b>":
            is_bold = True
        elif part == "</b>":
            is_bold = bold
        elif part:
            _set_run_font(paragraph.add_run(part), size=size, bold=is_bold, color=color)


def _para(container, text: str = "", size: float = 10, color: RGBColor = INK, bold: bool = False,
          space_after: float = 4, align=None):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    if text:
        _add_rich(p, text, size=size, color=color, bold=bold)
    return p


def _shade(cell, hex_fill: str) -> None:
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _table_borders(table, *, outer: str | None = None, inside_h: str | None = None, color: str = "D0D5DD",
                   size: int = 4) -> None:
    """Set table borders: outer box and/or horizontal lines between rows (no vertical lines)."""
    tblpr = table._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        style = outer if edge in ("top", "left", "bottom", "right") else (inside_h if edge == "insideH" else None)
        if style:
            el.set(qn("w:val"), style)
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tblpr.append(borders)


def _widths(table, widths_mm: list[float]) -> None:
    """Fixed layout with explicit widths on the grid and every cell (Word and LibreOffice agree)."""
    table.autofit = False
    tblpr = table._element.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
        for cell, w in zip(row.cells, widths_mm, strict=True):
            cell.width = Mm(w)
    for col, w in zip(table.columns, widths_mm, strict=True):
        col.width = Mm(w)


def _cell_margins(table, top: float = 1.2, bottom: float = 1.2, left: float = 1.5, right: float = 1.5) -> None:
    tblpr = table._element.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(int(val * 56.7)))  # mm -> twips
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def _cell_text(cell, text: str, size: float = 9.5, bold: bool = False, color: RGBColor = INK, align=None) -> None:
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    _add_rich(p, text, size=size, color=color, bold=bold)


def _hyperlink(paragraph, url: str, size: float = 9) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), FONT)
    rpr.append(fonts)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), "1849A9")
    rpr.append(col)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = url
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _field(paragraph, instr: str, size: float = 7.5) -> None:
    """Simple field such as PAGE."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    paragraph._p.append(fld)


def _writing_line(container, label: str, size: float = 10, label_width_mm: float = 32,
                  total_mm: float = CONTENT_W, space_after: float = 14):
    """'Label ________' using a right tab with an underline leader (prints as a writing line)."""
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.tab_stops.add_tab_stop(Mm(label_width_mm), WD_TAB_ALIGNMENT.LEFT)
    pf.tab_stops.add_tab_stop(Mm(total_mm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
    _set_run_font(p.add_run(f"{label}\t\t"), size=size, color=INK)
    return p


def _picture_border(inline_shape, color: str = "98A2B3", width_emu: int = 9525) -> None:
    """Thin outline on an inline picture (so white web pages stand out on white paper)."""
    sppr = inline_shape._inline.graphic.graphicData.pic.spPr
    ln = OxmlElement("a:ln")
    ln.set("w", str(width_emu))
    fill = OxmlElement("a:solidFill")
    clr = OxmlElement("a:srgbClr")
    clr.set("val", color)
    fill.append(clr)
    ln.append(fill)
    sppr.append(ln)
    # Reserve room for the outline, or Word clips its top edge against the line above.
    extent = inline_shape._inline.find(qn("wp:effectExtent"))
    if extent is None:
        extent = OxmlElement("wp:effectExtent")
        inline_shape._inline.insert(1, extent)
    for side in ("l", "t", "r", "b"):
        extent.set(side, str(width_emu))


def _box_choice(paragraph, label: str, size: float = 11) -> None:
    _set_run_font(paragraph.add_run(BOX + " "), name=SYMBOL_FONT, size=size + 3, color=INK)
    _set_run_font(paragraph.add_run(label + "      "), size=size, bold=True, color=INK)


# ---------------------------------------------------------------- document


def _setup(doc, meta: ReportMeta, T: dict[str, str]) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(MARGIN)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(18)
    sec.footer_distance = Mm(8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    normal.paragraph_format.space_after = Pt(4)

    footer_style = doc.styles["Footer"]
    footer_style.font.size = Pt(7.5)
    footer_style.font.color.rgb = MUTED
    footer = sec.footer
    fp = footer.paragraphs[0]
    tabs = fp.paragraph_format.tab_stops
    # The built-in Footer style has center/right tabs at 82.55 / 165.1 mm; clear them so the page
    # number sits at the right margin.
    for pos in (82.55, 165.1):
        tabs.add_tab_stop(Mm(pos), WD_TAB_ALIGNMENT.CLEAR)
    tabs.add_tab_stop(Mm(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
    _set_run_font(fp.add_run(f"{T['confidential']} · {meta.client_name} · {meta.case_id}\t{T['page']} "),
                  size=7.5, color=MUTED)
    _field(fp, "PAGE")

    cp = doc.core_properties
    cp.title = f"{T['title']} - {meta.case_id}"
    cp.author = meta.prepared_by or ""
    cp.subject = T["title"]


def _cover(doc, items: list[ReportItem], meta: ReportMeta, T: dict[str, str], lang: str) -> None:
    date_s = _fmt_date(meta.date, lang)
    _para(doc, "", space_after=36)
    _para(doc, T["title"], size=24, bold=True, space_after=10)

    info = [(T["prepared_for"], meta.client_name), (T["case"], meta.case_id), (T["date"], date_s)]
    if meta.prepared_by:
        info.append((T["prepared_by"], meta.prepared_by))
    t = doc.add_table(rows=len(info), cols=2)
    _table_borders(t)
    _cell_margins(t, left=0)
    for row, (k, v) in zip(t.rows, info, strict=True):
        _cell_text(row.cells[0], k, size=9, bold=True, color=MUTED)
        _cell_text(row.cells[1], v, size=10)
    _widths(t, [40, CONTENT_W - 40])

    _para(doc, "", space_after=10)
    _para(doc, T["intro1"].format(date=date_s), space_after=6)
    _para(doc, T["intro2"], space_after=6)
    _para(doc, PRINT_TEXT.get(lang, PRINT_TEXT["en"])["intro3"], space_after=12)

    _para(doc, T["summary"], size=12.5, bold=True, space_after=4)
    _para(doc, f"{T['total']}: <b>{len(items)}</b>", space_after=6)
    if not items:
        _para(doc, T["empty"])
    else:
        sev = [(s, sum(1 for i in items if i.get(H_SEVERITY) == s)) for s in SEVERITY_ORDER]
        sev = [(s, n) for s, n in sev if n]
        cats: dict[str, int] = {}
        for i in items:
            c = str(i.get(H_CATEGORY) or "Other")
            cats[c] = cats.get(c, 0) + 1
        cat = sorted(cats.items(), key=lambda x: -x[1])
        n_rows = 1 + max(len(sev), len(cat))
        t = doc.add_table(rows=n_rows, cols=5)
        _table_borders(t, inside_h="single", color="EAECF0")
        _cell_margins(t, left=0.5)
        headers = [T["by_severity"], T["count"], "", T["by_category"], T["count"]]
        for c, h in enumerate(headers):
            if h:
                _cell_text(t.rows[0].cells[c], h, size=8.5, bold=True, color=MUTED)
        for r, (s, n) in enumerate(sev, 1):
            _cell_text(t.rows[r].cells[0], tr(s, lang), bold=True,
                       color=RGBColor.from_string(SEVERITY_HEX[s]))
            _cell_text(t.rows[r].cells[1], str(n))
        for r, (c, n) in enumerate(cat, 1):
            _cell_text(t.rows[r].cells[3], tr(c, lang))
            _cell_text(t.rows[r].cells[4], str(n))
        _widths(t, [55, 20, 9, 70, 20])
    _para(doc, "", space_after=8)
    _para(doc, T["disclaimer"], size=8, color=MUTED)


LIST_BOOKMARK = "findings_list"


def _bookmark_name(item_id: int) -> str:
    return f"item_{item_id}"


def _add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    """Wrap the paragraph's content in a bookmark (target for the overview links)."""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    ppr = paragraph._p.pPr
    if ppr is not None:
        ppr.addnext(start)
    else:
        paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _internal_link(paragraph, text: str, anchor: str, size: float = 8, bold: bool = True) -> None:
    """Clickable jump to a bookmark in this document (survives Word's Save as PDF)."""
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    link.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), FONT)
    rpr.append(fonts)
    if bold:
        rpr.append(OxmlElement("w:b"))
    col = OxmlElement("w:color")
    col.set(qn("w:val"), "1849A9")
    rpr.append(col)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _overview(doc, items: list[ReportItem], T: dict[str, str], lang: str) -> None:
    """Compact list of all items after the cover; each number links to its item page."""
    p = _para(doc, T["list_title"], size=15, bold=True, space_after=4)
    p.paragraph_format.page_break_before = True
    _add_bookmark(p, LIST_BOOKMARK, 0)  # item bookmarks use ids >= 1
    _para(doc, T["list_intro"], size=8.5, color=MUTED, space_after=6)
    widths = [12, 40, 30, 44, 17, CONTENT_W - 143]
    headers = [T["no"], T["site"], T["site_type"], T["exposed"], T["severity"], T["action_short"]]
    t = doc.add_table(rows=len(items) + 1, cols=len(headers))
    _table_borders(t, inside_h="single", color="EAECF0")
    _cell_margins(t, top=0.6, bottom=0.6, left=0.8, right=0.8)
    for c, h in enumerate(headers):
        cell = t.rows[0].cells[c]
        _shade(cell, "F2F4F7")
        _cell_text(cell, h, size=7.5, bold=True, color=MUTED)
    for r, it in enumerate(items, 1):
        cells = t.rows[r].cells
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        _internal_link(p0, str(it.id), _bookmark_name(it.id))
        _cell_text(cells[1], _esc_markup(str(it.get(H_DOMAIN))), size=8)
        _cell_text(cells[2], _esc_markup(tr(it.get(H_CATEGORY), lang)), size=8)
        _cell_text(cells[3], _esc_markup(tr_list(it.get(H_DATA_TYPES), lang)), size=8)
        sev = str(it.get(H_SEVERITY) or "")
        _cell_text(cells[4], tr(sev, lang), size=8, bold=True,
                   color=RGBColor.from_string(SEVERITY_HEX[sev]) if sev in SEVERITY_HEX else INK)
        _cell_text(cells[5], _esc_markup(tr(it.get(H_METHOD), lang)), size=8)
    _widths(t, widths)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))  # repeat header on each page


def _item_page(doc, it: ReportItem, T: dict[str, str], lang: str) -> None:
    sev = str(it.get(H_SEVERITY) or "")
    head = doc.add_table(rows=1, cols=2)
    _table_borders(head)
    _cell_margins(head, left=0, right=0)
    left, right = head.rows[0].cells
    lp = left.paragraphs[0]
    lp.paragraph_format.page_break_before = True
    lp.paragraph_format.space_after = Pt(0)
    _set_run_font(lp.add_run(f"{T['item']} {it.id} · {it.get(H_DOMAIN)}"), size=15, bold=True, color=INK)
    _add_bookmark(lp, _bookmark_name(it.id), it.id)
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if sev:
        _shade(right, SEVERITY_HEX.get(sev, "475467"))
        _cell_text(right, tr(sev, lang), size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
                   align=WD_ALIGN_PARAGRAPH.CENTER)
    _widths(head, [CONTENT_W - 32, 32])
    _para(doc, "", size=4, space_after=4)

    rows: list[tuple[str, Any, bool]] = [
        (T["link"], it.get(H_URL), True),
        (T["site"], it.get(H_DOMAIN), False),
        (T["owner"], it.get(H_OWNER), False),
        (T["country"], it.get(H_JURISDICTION), False),
        (T["site_type"], tr(it.get(H_CATEGORY), lang), False),
        (T["content"], tr(it.get(H_CONTENT), lang), False),
        (T["exposed"], tr_list(it.get(H_DATA_TYPES), lang), False),
        (T["where"], tr_list(it.get(H_PLACES), lang), False),
    ]
    if it.get(H_MENTIONS) not in (None, ""):
        rows.append((T["mentions"], it.get(H_MENTIONS), False))
    rows += [
        (T["own"], tr(it.get(H_OWN), lang), False),
        (T["confidence"], tr(it.get(H_CONFIDENCE), lang), False),
        (T["action"], tr(it.get(H_METHOD), lang), False),
    ]
    if it.get(H_FALLBACK):
        rows.append((T["fallback"], tr(it.get(H_FALLBACK), lang), False))
    if it.get(H_CONTACT):
        rows.append((T["contact"], it.get(H_CONTACT), True))
    rows += [(T["status"], tr(it.get(H_STATUS), lang), False),
             (T["found"], _fmt_date(it.get(H_DATE_FOUND), lang), False)]

    t = doc.add_table(rows=len(rows), cols=2)
    _table_borders(t, inside_h="single", color="EAECF0")
    _cell_margins(t, left=0, top=1, bottom=1)
    for row, (label, value, is_link) in zip(t.rows, rows, strict=True):
        _cell_text(row.cells[0], label, size=8.5, bold=True, color=MUTED)
        if value in (None, ""):
            _cell_text(row.cells[1], T["none"])
        elif is_link:
            p = row.cells[1].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _hyperlink(p, str(value))
        else:
            _cell_text(row.cells[1], _esc_markup(str(value)))
    _widths(t, [44, CONTENT_W - 44])

    _para(doc, "", space_after=2)
    box = doc.add_table(rows=1, cols=1)
    _table_borders(box, outer="single", color="98A2B3", size=6)
    _cell_margins(box, top=2.5, bottom=2.5, left=3, right=3)
    cell = box.rows[0].cells[0]
    _shade(cell, "F2F4F7")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(6)
    _set_run_font(p.add_run(T["decision"] + "     "), size=11, bold=True, color=INK)
    _box_choice(p, T["remove"])
    _box_choice(p, T["keep"])
    _writing_line(cell, T["comment"], size=9, label_width_mm=18, total_mm=CONTENT_W - 8, space_after=2)
    _widths(box, [CONTENT_W])

    _para(doc, "", space_after=4)
    if it.evidence is not None:
        try:
            max_w, max_h = CONTENT_W - 2, 115.0
            buf, iw, ih, cropped = screenshot_jpeg(it.evidence, max_h / max_w)
            width = min(max_w, max_h * iw / ih)
            pp = doc.add_paragraph()
            pp.paragraph_format.space_after = Pt(2)
            pp.paragraph_format.keep_with_next = True
            shape = pp.add_run().add_picture(buf, width=Mm(width))
            _picture_border(shape)
            if cropped:
                _para(doc, T["cropped"], size=8, color=MUTED, space_after=0)
            return
        except Exception as exc:  # unreadable image: say so, don't fail the report
            log.warning("Could not embed screenshot for #%s: %s", it.id, type(exc).__name__)
    _para(doc, T["no_shot"], size=8, color=MUTED)


def _esc_markup(s: str) -> str:
    # Values are plain text; neutralize the <b> markup used by _add_rich.
    return s.replace("<b>", "‹b›").replace("</b>", "‹/b›")


def _authorization(doc, items: list[ReportItem], T: dict[str, str], lang: str) -> None:
    p = _para(doc, T["auth_title"], size=15, bold=True, space_after=8)
    p.paragraph_format.page_break_before = True
    _para(doc, T["auth_text"], space_after=10)
    if items:
        # The full list is the overview after the cover; don't repeat it here.
        p = _para(doc, f"{T['total']}: <b>{len(items)}</b> (", space_after=4)
        _internal_link(p, T["list_title"], LIST_BOOKMARK, size=10, bold=False)
        _set_run_font(p.add_run(")"), size=10, color=INK)
    else:
        _para(doc, T["empty"])
    _para(doc, "", space_after=24)
    for key in ("name", "signature", "sign_date"):
        _writing_line(doc, T[key], size=10, label_width_mm=32, total_mm=CONTENT_W * 0.8, space_after=22)


def build_docx(items: list[ReportItem], meta: ReportMeta, out_path: Path) -> Path:
    lang = meta.lang if meta.lang in TEXT else "en"
    T = TEXT[lang]
    doc = Document()
    _setup(doc, meta, T)
    _cover(doc, items, meta, T, lang)
    if items:
        _overview(doc, items, T, lang)
    for it in items:
        _item_page(doc, it, T, lang)
    _authorization(doc, items, T, lang)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def generate(
    tracker_path: Path,
    ident: ClientIdentifiers,
    out_path: Path,
    *,
    lang: str = "en",
    include_unverified: bool = False,
    order: str = "id",
    prepared_by: str = "",
    today: dt.date | None = None,
) -> tuple[Path, int]:
    tracker = Tracker(tracker_path)
    items = select_items(tracker, include_unverified, order)
    meta = ReportMeta(client_name=ident.full_name, case_id=ident.case_id, date=today or dt.date.today(),
                      prepared_by=prepared_by, lang=lang)
    build_docx(items, meta, out_path)
    return out_path, len(items)
