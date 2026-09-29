import datetime as dt

import openpyxl
import pytest

from footprint_scanner.excel_io import (
    H_CATEGORY,
    H_DATE_FOUND,
    H_DUE,
    H_ID,
    H_NOTES,
    H_RANK,
    H_SEVERITY,
    H_STATUS,
    H_URL,
    Tracker,
    TrackerError,
    timestamped_name,
)


def snapshot(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Findings"]
    return {
        "sheets": wb.sheetnames,
        "dv": sorted((dv.type, str(dv.sqref), dv.formula1) for dv in ws.data_validations.dataValidation),
        "cf": sorted((str(cf.sqref), r.type, tuple(r.formula or ()), r.dxfId, r.priority) for cf in ws.conditional_formatting for r in cf.rules),
        "dxf_count": len(wb._differential_styles.dxf),
        "formulas": {
            c: ws[c].value for c in ("A2", "W2", "X2", "A500", "W1000", "X1000")
        },
        "summary": [[c.value for c in row] for row in wb["Summary"].iter_rows()],
        "lists": [[c.value for c in row] for row in wb["Lists"].iter_rows()],
        "profile": [[c.value for c in row] for row in wb["Client Profile"].iter_rows()],
        "freeze": ws.freeze_panes,
        "filter": ws.auto_filter.ref,
        "profile_dv": [str(dv.sqref) for dv in wb["Client Profile"].data_validations.dataValidation],
    }


def test_allowed_values_read_from_validations(template_copy):
    t = Tracker(template_copy)
    assert t.allowed_for(H_CATEGORY)[0] == "Data broker / people search"
    assert "Re-appeared" in t.allowed_for(H_STATUS)
    assert t.allowed_for("Client's Own Account?") == ["Yes", "No", "Unknown"]
    assert "Google 'Results about you' / personal info form" in t.allowed_for("Deindex Fallback")
    with pytest.raises(TrackerError):
        t.validate(H_SEVERITY, "Severe")


def test_round_trip_preserves_everything(template_copy, tmp_path):
    before = snapshot(template_copy)
    t = Tracker(template_copy)
    assert t.clear_example_row()
    [row] = t.allocate_rows(1)
    assert row == 2
    t.set(row, H_URL, "https://example.com/a")
    t.set(row, H_DATE_FOUND, dt.date(2026, 9, 29))
    t.set(row, H_SEVERITY, "High")
    t.set(row, H_STATUS, "Not started")
    with pytest.raises(TrackerError):
        t.set(row, H_ID, 5)
    with pytest.raises(TrackerError):
        t.set(row, H_DUE, "x")
    out = t.save(tmp_path / "out.xlsx")
    after = snapshot(out)
    for k in before:
        if k != "summary":
            assert before[k] == after[k], k
    assert before["summary"] == after["summary"]

    ws = openpyxl.load_workbook(out)["Findings"]
    assert ws["C2"].value == "https://example.com/a"
    assert ws["A2"].value == "=IF(C2=\"\",\"\",ROW()-1)"
    assert ws["B2"].value == dt.datetime(2026, 9, 29) and ws["B2"].number_format == "yyyy\\-mm\\-dd"
    assert ws["AB2"].value is None  # example note cleared


def test_refuses_to_overwrite_input(template_copy):
    t = Tracker(template_copy)
    with pytest.raises(TrackerError):
        t.save(template_copy)


def test_columns_mapped_by_header_not_letter(template_copy, tmp_path):
    # Move 'Severity' and 'Notes' header text around: writing must follow the header.
    wb = openpyxl.load_workbook(template_copy)
    ws = wb["Findings"]
    ws.insert_cols(3)  # everything shifts one column to the right
    ws["C1"] = "Custom Column"
    wb.save(template_copy)
    t = Tracker(template_copy)
    t.set(5, H_NOTES, "hello")
    assert t.ws.cell(5, 29).value == "hello"  # AB -> AC


def test_existing_rows_and_example_detection(template_copy):
    t = Tracker(template_copy)
    assert t.is_example_row(2)
    assert t.existing_rows() == []  # example row doesn't count
    t.set(3, H_URL, "https://www.Example.com/x/?utm_source=a")
    [r] = t.existing_rows()
    assert r.row == 3 and r.id == 2 and r.key == "example.com/x"


def test_append_note_keeps_existing_text(template_copy):
    t = Tracker(template_copy)
    t.set(3, H_NOTES, "My own note")
    t.append_note(3, "[scanner] added")
    assert t.get(3, H_NOTES) == "My own note\n[scanner] added"


def test_growth_past_template_extends_formulas_and_ranges(template_copy, tmp_path):
    t = Tracker(template_copy)
    assert t.formula_extent == 1000
    t.clear_example_row()
    t.set(1000, H_URL, "https://x.com/last")
    rows = t.allocate_rows(3)
    assert rows == [1001, 1002, 1003]
    t.set(1003, H_URL, "https://x.com/new")
    t.set(1003, H_RANK, 4)
    out = t.save(tmp_path / "grown.xlsx")
    wb = openpyxl.load_workbook(out)
    ws = wb["Findings"]
    assert ws["A1003"].value == '=IF(C1003="","",ROW()-1)'
    assert ws["W1003"].value == '=IF(U1003="","",U1003+Summary!$C$4)'
    assert ws["A1003"].fill.fgColor.rgb == ws["A3"].fill.fgColor.rgb
    dvs = {dv.formula1: str(dv.sqref) for dv in ws.data_validations.dataValidation}
    assert dvs["Lists!$G$2:$G$10"] == "Y2:Y1003"
    cf_ranges = {str(cf.sqref) for cf in ws.conditional_formatting}
    assert "Q2:Q1003" in cf_ranges and "A2:AB1003" in cf_ranges
    assert ws.auto_filter.ref == "A1:AB1003"


def test_timestamped_name():
    from pathlib import Path

    when = dt.datetime(2026, 9, 29, 14, 30)
    assert timestamped_name(Path("Footprint_Removal_Tracker.xlsx"), when) == "Footprint_Removal_Tracker_2026-09-29_1430.xlsx"
    assert (
        timestamped_name(Path("Footprint_Removal_Tracker_2026-09-01_0900.xlsx"), when)
        == "Footprint_Removal_Tracker_2026-09-29_1430.xlsx"
    )


def excelify_validations(src, dst):
    """Rewrite cross-sheet list validations the way Excel saves them: in the x14 extLst."""
    import re
    import zipfile

    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet4.xml":
                xml = data.decode("utf-8")
                block = re.search(r"<dataValidations[^>]*>.*?</dataValidations>", xml, re.S).group(0)
                x14, keep = [], []
                for dv in re.findall(r"<dataValidation .*?</dataValidation>", block, re.S):
                    m = re.search(r"<formula1>(Lists![^<]+)</formula1>", dv)
                    if not m:
                        keep.append(dv)
                        continue
                    sqref = re.search(r'sqref="([^"]+)"', dv).group(1)
                    x14.append(
                        '<x14:dataValidation type="list" allowBlank="1" showErrorMessage="1" errorTitle="Invalid entry" '
                        'error="Pick a value from the list.">'
                        f"<x14:formula1><xm:f>{m.group(1)}</xm:f></x14:formula1><xm:sqref>{sqref}</xm:sqref>"
                        "</x14:dataValidation>"
                    )
                new_block = f'<dataValidations count="{len(keep)}">' + "".join(keep) + "</dataValidations>"
                xml = xml.replace(block, new_block)
                ext = (
                    '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
                    'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
                    f'<x14:dataValidations count="{len(x14)}" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
                    + "".join(x14) + "</x14:dataValidations></ext></extLst>"
                )
                xml = xml.replace("</worksheet>", ext + "</worksheet>")
                assert len(x14) == 7
                data = xml.encode("utf-8")
            zout.writestr(item, data)


def test_excel_extended_validations_are_restored(template_copy, tmp_path):
    excel_saved = tmp_path / "excel_saved.xlsx"
    excelify_validations(template_copy, excel_saved)
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        plain = openpyxl.load_workbook(excel_saved)
    assert len(plain["Findings"].data_validations.dataValidation) < 7  # what openpyxl alone would keep

    t = Tracker(excel_saved)
    assert t.allowed_for(H_STATUS)[-1] == "No action needed"
    assert t.allowed_for("Client's Own Account?") == ["Yes", "No", "Unknown"]
    assert t.allowed_for("Indexed in Google?") == ["Yes", "No", "Unknown"]
    with pytest.raises(TrackerError):
        t.validate(H_SEVERITY, "Severe")
    out = t.save(tmp_path / "out.xlsx")

    ws = openpyxl.load_workbook(out)["Findings"]
    dvs = {dv.formula1: str(dv.sqref) for dv in ws.data_validations.dataValidation}
    assert dvs["Lists!$G$2:$G$10"] == "Y2:Y1000"
    assert dvs["Lists!$C$2:$C$4"] == "I2:I1000 O2:P1000"
    assert len([f for f in dvs if f and f.startswith("Lists!")]) == 7
